import sys; sys.path.insert(0,'.')
from lib import *
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
IM='/tmp/claude-0/-home-user--/f2046ef5-a5f5-5769-8dca-a4878d4046a0/images/'
Image.open(IM+'10.webp').convert('RGB').save('intro/triptych.png'); Image.open(IM+'11.webp').convert('RGB').save('intro/mechanism.png')
prs=Presentation('/home/user/-/구두발표_콩vs로메인상추_v6.pptx')
sld=prs.slides._sldIdLst; ids=list(sld)
for i,x in enumerate(ids):
    if i not in (1,2): prs.part.drop_rel(x.rId); sld.remove(x)
S2,S3=prs.slides[0],prs.slides[1]
def byname(s,n,k=0): return [x for x in s.shapes if x.name==n][k]
def rm(sh): sh._element.getparent().remove(sh._element)
def tb(s,x,y,w,h,paras,size=11,color='333333',align=PP_ALIGN.CENTER,bold_first=False,anchor=MSO_ANCHOR.TOP):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=t.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(0.05); tf.margin_top=tf.margin_bottom=Inches(0.02)
    for i,p_ in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        runs=p_ if isinstance(p_,list) else [(p_,False)]
        for tx,b in runs:
            r=p.add_run(); r.text=tx; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=RGBColor.from_string(color)
    return t
# ---------------- slide 2
s=S2
for sh in list(s.shapes):
    if sh.name.startswith(('Card ','Header ','Text ')): rm(sh)
IMGX,IMGW=1.07,11.2; IMGY=2.15; IMGH=IMGW*520/1737
s.shapes.add_picture('intro/triptych.png',Inches(IMGX),Inches(IMGY),width=Inches(IMGW))
cent=[IMGX+IMGW*f for f in (0.1615,0.499,0.838)]
for k,nm in enumerate(('Flow 1','Flow 2','Flow 3')):
    b=byname(s,nm); b.left=Inches(cent[k]-1.65); b.width=Inches(3.3); b.top=Inches(1.3); b.height=Inches(0.78)
for k,nm in enumerate(('Arrow 1','Arrow 2')):
    a=byname(s,nm); a.left=Inches((cent[k]+cent[k+1])/2-0.2); a.width=Inches(0.4); a.top=Inches(1.3+0.39-0.15)
caps=[[[('근권 산소 결핍 → 기공 닫힘',True)],[('엽 내 CO',False),('2','sub'),(' 공급 감소, 회복 지연',False)]],
      [[('토양 수분 부족 → 엽 수분 저하',True)],[('증산·광합성 감소',False)]],
      [[('침수 회복 전 건조 노출',True)],[('피해 누적 또는 순화 가능',False)]]]
for k,c in enumerate(caps):
    t=s.shapes.add_textbox(Inches(cent[k]-1.75),Inches(IMGY+IMGH+0.08),Inches(3.5),Inches(0.7)); tf=t.text_frame; tf.word_wrap=True
    for i,line in enumerate(c):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.CENTER
        for tx,kind in line:
            r=p.add_run(); r.text=tx; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string('333333')
            if i==0: r.font.bold=True
            if kind=='sub': r._r.get_or_add_rPr().set('baseline','-25000')
set_para(byname(s,'Footnote').text_frame.paragraphs[0],'극한기상의 빈도·강도 증가와 복합재해 증가 (기상청, 2025; 김용준 등, 2025). 2025년 7월 중순 남부지방 집중호우(5일 누적 강수량 500 mm 이상) 후 폭염이 지속됨 (KREI, 2025). 이미지: 생성형 AI로 제작한 개념도.')
byname(s,'Footnote').height=Inches(0.34)
s.notes_slide.notes_text_frame.text='(약 35초) 최근에는 집중호우와 폭염·가뭄이 짧은 간격으로 반복되면서 한 작물이 한 생육기 안에서 침수와 건조를 연이어 겪는 복합재해가 늘고 있습니다. 2025년 7월 중순에도 남부지방에 집중호우가 내린 뒤 폭염이 이어졌습니다. 왼쪽 그림처럼 침수는 뿌리의 산소 결핍으로 기공을 닫게 하고 배수 후에도 회복에 시간이 걸립니다. 가운데 그림처럼 건조는 토양 수분 부족으로 엽의 수분 상태를 떨어뜨려 증산과 광합성을 줄입니다. 오른쪽 그림처럼 두 재해가 이어지면 피해가 누적될 수도, 앞선 스트레스가 순화 효과를 줄 수도 있는데, 작물과 생육 단계별로 어떻게 다른지 비교한 연구는 충분하지 않습니다.'
# ---------------- slide 3
s=S3
for nm in ('Step 1','Step 2','Step 4','Step 5'):
    pass
steps=[byname(s,'Step %d'%i) for i in (1,2,4,5)]
IW=7.0; IH=IW*619/1108
s.shapes.add_picture('intro/mechanism.png',Inches(0.92),Inches(1.3),width=Inches(IW))
ob=byname(s,'Objective'); ob.top=Inches(5.3); ob.height=Inches(0.95)
rows=[('① 기공','g_{sw}','수분 스트레스에 가장 먼저 반응해 닫히고, 엽 내 CO2 공급이 감소'),
      ('② 명반응','Φ_{PSII}','틸라코이드 막 PSII 전자전달이 저해되어 광화학 효율 저하'),
      ('③ 광합성 능력','V_{cmax} · J_{max} · A_{max}','루비스코 카르복실화와 RuBP 재생 능력 저하 (②와 병렬로 평가)'),
      ('④ 생육·수량','수량 · 지상부 건물중','①~③의 영향이 누적되어 생육·수량 감소로 이어지는지 평가')]
for k,(st,(a,b,c)) in enumerate(zip(steps,rows)):
    ps=st.text_frame.paragraphs; set_para(ps[0],a); set_para(ps[1],b)
    st.left=Inches(8.1); st.top=Inches(1.3+k*0.98); st.width=Inches(1.85); st.height=Inches(0.86)
    t=tb(s,10.05,1.3+k*0.98,2.45,0.86,[c],size=10.5,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE)
    for r in t.text_frame.paragraphs[0].runs:
        pass
    # subscript for CO2
    if k==0:
        p=t.text_frame.paragraphs[0]; [p._p.remove(r._r) for r in list(p.runs)]
        for tx,sub in [('수분 스트레스에 가장 먼저 반응해 닫히고, 엽 내 CO',False),('2',True),(' 공급이 감소',False)]:
            r=p.add_run(); r.text=tx; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string('333333')
            if sub: r._r.get_or_add_rPr().set('baseline','-25000')
s.notes_slide.notes_text_frame.text='(약 30초) 이 연구는 수분 재해의 영향을 단계별로 나누어 보았습니다. 잎에서는 수분 스트레스에 가장 먼저 기공이 닫혀 이산화탄소 공급이 줄고, 그 뒤 두 경로가 병렬로 영향을 받습니다. 하나는 틸라코이드 막의 PSII 전자전달, 즉 명반응이고, 다른 하나는 루비스코 카르복실화와 재생 능력, 즉 광합성 능력입니다. 이 세 단계의 영향이 누적되어 생육과 수량 감소로 이어지는지를 콩과 로메인 상추에 침수, 건조, 복합 처리를 적용하여 비교했습니다.'
for _s in prs.slides: fix_timing(_s)
prs.save('intro/서론_2장.pptx'); print('saved',len(prs.slides))
