import sys,re,copy; sys.path.insert(0,'.')
from lib import *
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import pandas as pd, numpy as np
from PIL import Image
R='/home/user/-/'; AN='an/'
src=open('oral_b.py').read()
for fn in ('bottom_border','compact_table','drop_col','set_widths'):
    m=re.search(r'\ndef '+fn+r'\(.*?(?=\ndef |\n# =====|\nMM=)',src,flags=re.S); exec(m.group(0))
prs=Presentation(R+'구두발표_콩vs로메인상추_v8.pptx')
SL=list(prs.slides)               # slide objects by original index
S=lambda n: SL[n-1]
def byname(s,n,k=0): return [x for x in s.shapes if x.name==n][k]
def allnames(s,n): return [x for x in s.shapes if x.name==n]
def rm(sh): sh._element.getparent().remove(sh._element)
def img(p): return Image.open(p).convert('RGB')
def notes(s,t): s.notes_slide.notes_text_frame.text=t
def conc(s,t): set_para(byname(s,'Conclusion Box').text_frame.paragraphs[0],t)
def foot(s,t):
    try: set_para(byname(s,'Footnote').text_frame.paragraphs[0],t)
    except IndexError: pass
PNS='{http://schemas.openxmlformats.org/presentationml/2006/main}'
ANS='{http://schemas.openxmlformats.org/drawingml/2006/main}'
def newids(s,els):
    mx=max(int(e.get('id')) for e in s.shapes._spTree.iter(PNS+'cNvPr'))
    for el in els:
        for c in el.iter(PNS+'cNvPr'): mx+=1; c.set('id',str(mx))
def copy_el(dst,el):
    e=copy.deepcopy(el); dst.shapes._spTree.append(e); newids(dst,[e]); return e
def shp(s,el): return [x for x in s.shapes if x._element is el][0]
def add_pic(s,path,x,y,w):
    p=s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w)); return p
def tbl(s,k=0): return [x for x in s.shapes if getattr(x,'has_table',False) and x.has_table][k]
def instr_boxes(s):   # user's instruction boxes: autoshapes named '직사각형 N'
    return [x for x in s.shapes if x.name.startswith('직사각형')]
def rewrite_runs(p,parts,size=None):
    rs=p._p.findall(ANS+'r'); base=rs[0]
    for r in rs[1:]: p._p.remove(r)
    prev=None
    for i,(t,b) in enumerate(parts):
        r=base if i==0 else copy.deepcopy(base)
        if i>0: prev.addnext(r)
        r.find(ANS+'t').text=t; rpr=r.find(ANS+'rPr')
        if rpr is not None:
            if 'baseline' in rpr.attrib: del rpr.attrib['baseline']
            rpr.set('b','1' if b else '0')
            if size: rpr.set('sz',str(int(size*100)))
        prev=r
MM=pd.read_csv(AN+'intermediate/mm_phase.csv')
def fill_mm(t,ind,fmt,name=None):
    while len(t.rows)>5: delete_row(t,len(t.rows)-1)
    for c_i,h in enumerate([name or 'Treatment','Soybean P1','Soybean P2','Lettuce P1','Lettuce P2']): set_cell(t.cell(0,c_i),h)
    for r_i,(tn,tr) in enumerate([('Control','ETc'),('Drought','ETc50'),('Waterlogging','침수ETc'),('Combined','침수ETc50')],start=1):
        set_cell(t.cell(r_i,0),tn)
        for c_i,(cr,pe) in enumerate([('bean','P1'),('bean','P2'),('lettuce','P1'),('lettuce','P2')],start=1):
            q=MM[(MM.지표==ind)&(MM.작물==cr)&(MM.구간==pe)&(MM.처리==tr)]
            set_cell(t.cell(r_i,c_i),'–' if q.empty else (fmt.format(q.보정평균.iloc[0])+' '+q.문자.iloc[0]))
def replace_pic_keep(s,name,path,nth=0):
    pic=byname(s,name,nth); return replace_picture(s,pic,img(path),valign='top')
# ================= global: Romaine -> Lettuce
for s in SL:
    for sh in s.shapes:
        tfs=[]
        if sh.has_text_frame: tfs.append(sh.text_frame)
        if getattr(sh,'has_table',False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells: tfs.append(c.text_frame)
        for tf in tfs:
            for p in tf.paragraphs:
                for r in p.runs:
                    if 'Romaine' in r.text: r.text=r.text.replace('Romaine Lettuce','Lettuce').replace('Romaine','Lettuce')
# ================= slide 3
s=S(3)
for b in instr_boxes(s): rm(b)
ps=byname(s,'Objective').text_frame.paragraphs
rewrite_runs(ps[1],[('잎 생리 지표 중 ',False),('어느 지표가 먼저 손상되는지',True),(', ',False),('어느 지표가 회복되는지',True),(', ',False),('생육·수량 감소로 이어지는지',True),(' 비교',False)])
# ================= slide 4
s=S(4)
for b in instr_boxes(s): rm(b)
t=s.shapes.add_textbox(Inches(6.7),Inches(5.3),Inches(5.8),Inches(1.0)); tf=t.text_frame; tf.word_wrap=True
for i,(lead,body) in enumerate([('관수 : ','전날 증산한 만큼을 다음날 관수 (건조·복합은 50%)'),('침수 : ','포트를 물에 담가 72시간 침지')]):
    p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
    for tx,b in ((lead,True),(body,False)):
        r=p.add_run(); r.text=tx; r.font.size=Pt(13); r.font.bold=b; r.font.color.rgb=RGBColor.from_string('333333')
notes(s,'(약 20초) 공시 작물은 상추와 콩 두 종이며, 각각 대조, 건조, 침수, 복합 네 처리구를 설정했습니다. 관수는 식물체가 전날 증산한 만큼 다음날 관수했고, 건조와 복합 처리는 그 50%만 관수했습니다. 침수는 포트를 물에 담가 72시간, 즉 3일간 처리했고 복합 처리는 침수 직후 건조 처리를 이어서 했습니다.')
# ================= slide 6
s=S(6)
for b in instr_boxes(s): rm(b)
tops=[1.35,2.50,3.85,4.95]; hs=[1.05,1.25,0.95,0.78]
bgs=allnames(s,'Row bg'); chips=allnames(s,'Row chip'); inds=allnames(s,'Row ind'); mets=allnames(s,'Row method')
for k in range(4):
    for g in (bgs,chips,inds,mets): g[k].top=Inches(tops[k]); g[k].height=Inches(hs[k])
    inds[k].width=Inches(3.0); mets[k].left=Inches(6.6); mets[k].width=Inches(5.75)
def setm(sh,lines,size=10.5):
    set_paras(sh,lines)
    for p in sh.text_frame.paragraphs:
        for r in p.runs: r.font.size=Pt(size)
setm(mets[0],['LI-600 포로미터, 처리 개시 후 반복 측정','온도 보정 gsw (Rizzo & Bailey, 2026)'],11)
setm(mets[1],['LI-600 명적응 형광 (Fs, Fm′) + 암적응 Fo, Fm','ΦNO = 1/(NPQ+1+qL(Fm/Fo−1)), ΦNPQ = 1−ΦPSII−ΦNO (Kramer et al., 2004)','보정: y = μ + 처리 + β_{1}·PAR + β_{2}·기온 + 측정일(임의효과) (Ali et al., 2025; Gao et al., 2024)'],10)
setm(mets[2],['LI-6800, A-Ci (FvCB) → Vcmax, Jmax, Rd','A-PPFD (NRH 모형) → Amax, α'],11)
setm(mets[3],['생육 조사 (상추 DAT 16, 콩 DAT 55)'],11)
set_para(byname(s,'ToolSection').text_frame.paragraphs[0],'매개변수 산출 ▸ A-Ci: R·plantecophys / A-PPFD: NRH 모형')
for p in byname(s,'ToolSection').text_frame.paragraphs:
    for r in p.runs: r.font.size=Pt(11)
notes(s,'(약 30초) 측정 항목은 네 단계입니다. 기공은 온도 보정한 기공전도도로 보았습니다. 명반응은 LI-600 형광으로 광계 II 양자수율 ΦPSII를 얻고, 암적응 Fo와 Fm을 이용해 NPQ, ΦNO, ΦNPQ를 계산했습니다. 형광 지표는 측정 당시의 PAR과 기온을 양적 공변량으로, 측정일을 임의효과로 넣은 모형으로 보정했습니다. 광합성 능력은 A-Ci 곡선에서 Vcmax, Jmax, Rd를, 광-광합성 곡선에 비직각쌍곡선 모형을 적합해 Amax와 α를 구했고, 마지막으로 생육과 수량을 조사했습니다.')
# ================= slide 7
s=S(7)
for b in instr_boxes(s): rm(b)
for nm in ('Stat box','Stat value','Stat label'):
    for x in allnames(s,nm): rm(x)
replace_pic_keep(s,'Picture 26',AN+'figs4/env.png')
e1=pd.read_csv(R+'data/19_환경_밭콩_일별.csv'); e2=pd.read_csv(R+'로메인data/25_환경_일별.csv')
a=e1[(e1.DAT>=0)&(e1.DAT<=16)]; b=e2[(e2.DAT>=0)&(e2.DAT<=16)&(e2['결측일(자료<1400)'].astype(str)!='True')]
rows=[('평균기온 (°C)',b.평균기온.mean(),a.평균기온.mean(),'{:.1f}'),('DLI (mol·m⁻²·d⁻¹)',b['DLI_추정(mol m-2 d-1)'].mean(),a.DLI_mol.mean(),'{:.1f}'),('최고기온 (°C)',b.최고기온.mean(),a.최고기온.mean(),'{:.1f}'),('최저기온 (°C)',b.최저기온.mean(),a.최저기온.mean(),'{:.1f}'),('평균 상대습도 (%)',b.평균습도.mean(),a.평균습도.mean(),'{:.1f}')]
ts=tbl(s); tt=ts.table
while len(tt.rows)<6: tt._tbl.append(copy.deepcopy(tt._tbl.tr_lst[-1]))
for c_i,h in enumerate(['처리 기간 (DAT 0~16) 평균','상추','콩']): set_cell(tt.cell(0,c_i),h)
for r_i,(nm,x1,x2,f) in enumerate(rows,start=1):
    set_cell(tt.cell(r_i,0),nm); set_cell(tt.cell(r_i,1),f.format(x1)); set_cell(tt.cell(r_i,2),f.format(x2))
set_widths(ts,[3.0,2.0,2.0],0.28); ts.left=Inches(0.9); ts.top=Inches(4.55)
t=s.shapes.add_textbox(Inches(8.6),Inches(3.8),Inches(3.8),Inches(0.6)); tf=t.text_frame; tf.word_wrap=True
p=tf.paragraphs[0]
for tx,bd in (('재배 장소 : ',True),('전북대학교 부속농장 온실',False)):
    r=p.add_run(); r.text=tx; r.font.size=Pt(13); r.font.bold=bd; r.font.color.rgb=RGBColor.from_string('333333')
foot(s,'밭콩: CR1000X 기상 센서(8/7~10/6), 상추: 26번 온실 환경 데이터(CR300; 8/22~10/23, 10/13·10/18 결측일 제외). 상추 DLI = 일사량(MJ) × 2.02 mol/MJ 추정. 점선 = 침수 시작일. 표: DAT 0~16 일별 값의 평균.')
notes(s,'(약 25초) 재배 환경입니다. 두 실험 모두 전북대학교 부속농장 온실에서 진행했고, 그래프는 측정된 전체 재배 기간의 기온과 일적산광량입니다. 처리 기간인 DAT 0~16의 평균 기온은 콩이 28.4도, 상추가 18.2도로 약 10도 차이가 났고 일적산광량, 최고·최저 기온과 상대습도도 달랐습니다. 같은 처리라도 스트레스 강도가 다를 수 있어서 두 작물은 병렬 사례로 비교하겠습니다.')
# ================= slide 8: nav -> 결과, picture
s=S(8)
for b in instr_boxes(s): rm(b)
for nm in ('Nav 서론','Nav 재료 및 방법','Nav 결과','Nav 고찰','Nav 결론'):
    a=byname(s,nm); b=byname(S(9),nm)
    for pa,pb in zip(a.text_frame.paragraphs,b.text_frame.paragraphs):
        for ra,rb in zip(pa.runs,pb.runs):
            old=ra._r.find(ANS+'rPr'); new=copy.deepcopy(rb._r.find(ANS+'rPr'))
            if old is not None: ra._r.replace(old,new)
replace_pic_keep(s,'Picture 26',AN+'figs3/irr_cum.png')
# ================= slide 9: equal-axes pair
s=S(9)
for b in instr_boxes(s): rm(b)
for nm in ('Picture 19','Picture 20'): rm(byname(s,nm))
W=6.4; h=W*2.55/7.5
add_pic(s,AN+'figs4/gsw_ts.png',0.5,1.25,W); add_pic(s,AN+'figs4/gsw_rec.png',0.5,1.25+h+0.12,W)
# ================= slide 10 (merged light reactions) ; slide 11 removed later
s=S(10)
for nm in ('Picture 13','TextBox 10'): rm(byname(s,nm))
base_t=tbl(s)
set_para(s.shapes.title.text_frame.paragraphs[0],'LI-600 : 명반응 (Φ_{PSII}, Φ_{NPQ}, Φ_{NO})')
add_pic(s,AN+'figs4/phi3.png',0.5,1.2,6.3)
t1=base_t; t1.left=Inches(7.05); t1.top=Inches(1.5); fill_mm(t1.table,'phi','{:.3f}','Φ_{PSII}'); set_widths(t1,[1.4,1.0,1.0,1.0,1.0],0.27)
for k,(ind,nm) in enumerate([('phinpq','Φ_{NPQ}'),('phino','Φ_{NO}')],start=1):
    e=copy_el(s,t1._element); tx=shp(s,e); tx.top=Inches(1.5+k*1.6); fill_mm(tx.table,ind,'{:.3f}',nm); set_widths(tx,[1.4,1.0,1.0,1.0,1.0],0.27)
conc(s,'콩은 ΦPSII가 낮고 ΦNPQ·ΦNO가 높아 명반응에 타격을 입은 것으로 추정, 상추는 처리 간 차이 없음')
foot(s,'Φ_{NPQ}=1−Φ_{PSII}−Φ_{NO}, Φ_{NO}=1/(NPQ+1+qL(Fm/Fo−1)) (Kramer et al., 2004). 표·그림: PAR·기온 공변량 보정 평균(혼합모형, 측정일 임의효과; Ali et al., 2025; Gao et al., 2024), Tukey 문자. 그림: 음영 ± SE, 빗금 침수 기간, 점선 P1/P2 경계, * 당일 P<0.05(Holm).')
notes(s,'(약 55초) 명반응입니다. 흡수한 빛에너지는 광화학에 쓰이는 ΦPSII, 조절된 열소산 ΦNPQ, 조절되지 않는 소산 ΦNO로 나뉘고 세 값의 합은 1입니다. 콩은 P1부터 침수구의 ΦPSII가 대조구보다 낮았고, P2에서는 건조구가 가장 낮았으며 복합구도 낮았습니다. ΦNPQ는 P1에서 침수구가, P2에서 건조구가 높았고 복합구는 높은 경향이었습니다. ΦNO는 P2에서 대조구만 낮고 건조, 침수, 복합 처리가 모두 높았는데, ΦNO는 광계 손상의 위험을 나타내므로 콩은 처리 전반에서 명반응이 타격을 입은 것으로 추정됩니다. 반면 상추는 세 지표 모두 전 구간에서 처리 간 차이가 없었습니다.')
# ================= slide 12 (photosynthesis) picture
s=S(12); replace_pic_keep(s,'Picture 13',AN+'figs3/params.png')
# ================= slide 13 (yield): transposed tables
s=S(13)
for b in instr_boxes(s): rm(b)
replace_pic_keep(s,'Picture 26',AN+'figs3/yield.png')
yt=pd.read_csv(AN+'intermediate/yield_table.csv')
old=tbl(s); src_el=old._element
def pc(v,dd,let,fm):
    return fm(v)+(f' {let}' if isinstance(let,str) and let else '')+('' if dd is None else f' ({"−" if dd<0 else "+"}{int(abs(dd)+0.5)}%)')
def build(idxs,heads,name,left,w):
    e=copy_el(s,src_el); tx=shp(s,e); tt=tx.table
    while len(tt.rows)>5: delete_row(tt,len(tt.rows)-1)
    while len(tt.columns)>1+len(idxs): drop_col(tt,len(tt.columns)-1)
    set_cell(tt.cell(0,0),name)
    for c_i,h in enumerate(heads,start=1): set_cell(tt.cell(0,c_i),h)
    for r_i,(tn,k) in enumerate([('Control',0),('Drought',1),('Waterlog.',2),('Combined',3)],start=1):
        set_cell(tt.cell(r_i,0),tn)
        for c_i,i in enumerate(idxs,start=1):
            r=yt.iloc[i]; vals=[r.ETc,r.ETc50,r.침수ETc,r.침수ETc50]; ds=[None,r.d_ETc50,r.d_침수,r.d_복합]; lets=[r.l0,r.l1,r.l2,r.l3]
            fm=(lambda v:f'{v:.1f}') if vals[0]<20 else (lambda v:f'{v:.0f}')
            set_cell(tt.cell(r_i,c_i),pc(vals[k],ds[k],lets[k],fm))
    set_widths(tx,w,0.3); tx.left=Inches(left); tx.top=Inches(3.75)
build([0,1,2,3],['Seed FW (g)','Seed DW (g)','Seed no.','Pod no.'],'Soybean DAT 55',0.5,[1.2,1.05,1.05,0.95,0.95])
build([4,5],['Shoot FW (g)','Shoot DW (g)'],'Lettuce DAT 16',6.05,[1.15,1.15,1.15])
rm(old)
foot(s,'표: 평균 (대조구 대비 %), 건조·침수·복합 순. 상추 10/23(DAT 16) 처리당 8~9주, 콩 10/6(DAT 55, R6~R7) 파괴조사 n=6~9. 그림: 평균 ± SE, 문자 Tukey(ANOVA P<0.05일 때만).')
# ================= slides 15/16 -> 콩/상추 (총정리 + 고찰)
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
fig=plt.figure(figsize=(12.2,3.9)); gs=fig.add_gridspec(1,4,wspace=0.3,left=0.055,right=0.995,top=0.9,bottom=0.15,width_ratios=[1,1.05,0.95,0.95])
l,r=[np.array(x) for x in gs.get_grid_positions(fig)[2:4]]; cen=(l+r)/2; plt.close(fig)
def recap_slide(s,crop,ttl,cc_,ft,nt,labels):
    set_para(s.shapes.title.text_frame.paragraphs[0],ttl)
    rm(byname(s,'Picture 64')); p2=add_pic(s,AN+f'figs4/recap2_{crop}.png',0.55,2.35,12.2)
    s.shapes._spTree.remove(p2._element); s.shapes._spTree.insert(2,p2._element)
    rects=sorted([x for x in s.shapes if x.name.startswith('Rounded Rectangle')],key=lambda r:r.left)
    for r_,c_ in zip(rects,cen): r_.left=Inches(0.55+12.2*c_-1.35)
    set_para(rects[1].text_frame.paragraphs[1],'Φ_{PSII} · Φ_{NPQ} · Φ_{NO} 비율')
    arrows=sorted([x for x in s.shapes if x.name.startswith('Right Arrow')],key=lambda r:r.left)
    for k,a in enumerate(arrows): a.left=Inches((0.55+12.2*cen[k]+0.55+12.2*cen[k+1])/2-0.25)
    tags=sorted(instr_boxes(s),key=lambda r:r.left)
    for tg,c_,lab in zip(tags,cen,labels):
        tg.width=Inches(1.5); tg.height=Inches(0.3); tg.left=Inches(0.55+12.2*c_-0.75); tg.top=Inches(2.0)
        pp=tg.text_frame.paragraphs[0]; set_para(pp,lab)
    conc(s,cc_); foot(s,ft); notes(s,nt)
recap_slide(S(15),'bean','콩 : 지표별 반응 총정리 및 고찰','콩은 생리 지표에서 차이가 있었고, 수량에는 감소하는 경향이 보였지만 큰 차이가 없었다',
 '① 건조·복합 회색 / ③ 침수 회색 (영향이 큰 처리만 색 강조), ② 흡수 에너지 분배(%, 혼합모형 보정 평균), ③ 처리 전 제외. Tan et al. 2018; Haverroth et al. 2024; Buckley 2019; Yeung et al. 2018; Sasidharan et al. 2018 (참고문헌: 보충 자료)',
 '(약 75초) 지표를 연결해서 콩부터 정리하겠습니다. 침수 직후 뿌리가 저산소 상태가 되면 수분 흡수가 떨어져 기공이 닫히고 CO2 공급이 줄어드는데, 콩은 이 회복에 2~3주가 걸렸고 건조와 복합은 낮게 유지되었습니다. 기공이 닫혀도 빛은 계속 흡수되므로 두 번째 그래프처럼 광화학에 쓰이는 ΦPSII는 줄고 ΦNPQ와 ΦNO의 비율이 늘었습니다. 건조구는 ΦNPQ를 높여 열로 보호하려 했지만 ΦNO도 높아 광계에 부담이 있었고, 세 번째 그래프처럼 Vcmax와 Jmax도 건조구에서 낮았습니다. 그런데도 수확량에 큰 차이가 없었던 것은, 대조구와 침수구의 관수량 차이가 작아 3일 침수의 타격이 작았고, 처리 후 회복 기간 동안 종자 크기로 보상되었을 수 있으며, 조기 수확과 적은 반복으로 차이를 검출하기 어려웠기 때문으로 해석합니다. 다만 기공, 명반응, 광합성 능력이 모두 나빴던 복합구에서 수확량이 가장 낮았다는 점은 생리 지표와 같은 방향입니다.',
 ['건조, 복합 ↓','건조, 복합 ↑','건조, 복합 ↓','복합 ↓'])
recap_slide(S(16),'lettuce','상추 : 지표별 반응 총정리 및 고찰','상추는 회복이 빨랐지만 침수 기간 동안의 손실이 유지됨',
 '① 건조·복합 회색 / ③ 건조·침수 회색 (영향이 큰 처리만 색 강조), ② 흡수 에너지 분배(%, 혼합모형 보정 평균), ③ 처리 전 제외. Körner 2015; Tardieu et al. 2018; Haverroth et al. 2024 (참고문헌: 보충 자료)',
 '(약 70초) 다음은 상추입니다. 첫 번째 그래프처럼 침수 중 기공이 대조구의 10%까지 떨어졌지만 배수 후 2~3일 만에 회복했고, 두 번째 그래프처럼 흡수한 에너지의 분배는 처리 간 차이가 없어 명반응은 영향을 받지 않았습니다. 광합성 능력은 복합구에서만 Vcmax와 Jmax가 낮았습니다. 그런데도 생육은 크게 줄었는데, 대조구와 침수구의 관수량 차이가 컸다는 점에서 침수가 생육에 직접적인 타격을 주었을 가능성이 있습니다. 상추는 영양생장기에 잎이 수확부라서 침수 3일 동안 멈춘 탄소 고정과 잎 생장의 손실이 회복 후에도 만회되지 않고 남은 것으로 해석됩니다. 식물의 생장은 탄소 공급보다 조직 성장이 먼저 제한되는 경우가 많아 기공이 회복되어도 이미 멈춘 생장은 따라잡기 어렵습니다. 복합구는 여기에 건조 후반의 기공 감소와 Vcmax, Jmax 저하가 더해져 가장 크게 감소했습니다.',
 ['건조 ↓','차이X','복합 ↓','복합 ↓'])
# tag labels (after tags were sorted in recap_slide, texts already set by old shapes; set explicitly)
for s,labs in ((S(15),['건조, 복합 ↓','건조, 복합 ↑','건조, 복합 ↓','복합 ↓']),(S(16),['건조 ↓','차이X','복합 ↓','복합 ↓'])):
    for tg,lab in zip(sorted(instr_boxes(s),key=lambda r:r.left),labs): set_para(tg.text_frame.paragraphs[0],lab)
# ================= appendix pictures
for n,fig in ((21,'figs3/aq.png'),(22,'figs2/ts_fs.png'),(23,'figs2/ts_etr.png')):
    replace_pic_keep(S(n),'Picture 13',AN+fig)
# ================= delete slides 11 and 14
sld=prs.slides._sldIdLst; ids=list(sld)
for n in (14,11):
    prs.part.drop_rel(ids[n-1].rId); sld.remove(ids[n-1])
# borders / compact / timing
for s in prs.slides:
    for sh in s.shapes:
        if getattr(sh,'has_table',False) and sh.has_table: compact_table(sh.table); bottom_border(sh.table)
    fix_timing(s)
prs.save('oral/oral_v9.pptx'); print('saved',len(prs.slides))
