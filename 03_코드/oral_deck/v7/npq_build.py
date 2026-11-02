import sys,re,copy; sys.path.insert(0,'.')
from lib import *
from pptx.util import Pt
import pandas as pd
from PIL import Image
src=open('oral_b.py').read()
for fn in ('bottom_border','compact_table'):
    m=re.search(r'\ndef '+fn+r'\(.*?(?=\ndef |\n# =====|\ndef pct)',src,flags=re.S); exec(m.group(0))
prs=Presentation('/home/user/-/구두발표_콩vs로메인상추_v6.pptx')
sld=prs.slides._sldIdLst; ids=list(sld)
idx=[i for i,s in enumerate(prs.slides) if s.shapes.title is not None and s.shapes.title.text_frame.text.startswith('LI-600 : 정상 상태 형광')][0]
for i,x in enumerate(ids):
    if i!=idx: prs.part.drop_rel(x.rId); sld.remove(x)
base=prs.slides[0]
def dup(srcs):
    new=prs.slides.add_slide(srcs.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    for el in srcs.shapes._spTree:
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        e2=copy.deepcopy(el)
        for b in e2.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            rid=b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if rid:
                tp=srcs.part.rels[rid].target_part
                b.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',new.part.relate_to(tp,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'))
        new.shapes._spTree.append(e2)
    return new
slides=[base,dup(base),dup(base)]
def byname(s,n,k=0): return [x for x in s.shapes if x.name==n][k]
def conc(s,t): set_para(byname(s,'Conclusion Box').text_frame.paragraphs[0],t)
def foot(s,t): set_para(byname(s,'Footnote').text_frame.paragraphs[0],t)
P=pd.read_csv('an/intermediate/mm_phase.csv'); P=P[P.구간.isin(['P1','P2'])]
INDMAP={'NPQ':'npq','PhiNO':'phino','PhiNPQ':'phinpq'}
def pv(crop,per,ind,tr):
    q=P[(P.작물==crop)&(P.구간==per)&(P.지표==INDMAP[ind])&(P.처리==tr)]
    return '–' if q.empty else f'{q.보정평균.iloc[0]:.3f} {q.문자.iloc[0]}'
SPEC=[('NPQ','ts_npq.png','LI-600 : 비광화학적 소광 (NPQ)',
  ['콩: P1은 처리 간 차이가 없었고, P2는 건조구가 침수구보다 높았음','로메인: P1·P2 모두 처리 간 차이 없음'],
  '콩은 P2에서 건조구의 NPQ가 침수구보다 높고, 로메인은 처리 간 차이 없음',
  '(약 25초) 비광화학적 소광 NPQ입니다. 콩은 P1에서는 처리 간 차이가 없었고 P2에서는 건조구가 침수구보다 높았습니다. 로메인 상추는 P1과 P2 모두 처리 간 차이가 없었습니다. 건조 스트레스에서 콩이 과잉 에너지를 열로 더 많이 소산했다는 뜻입니다.'),
 ('PhiNO','ts_phino.png','LI-600 : 비조절 에너지 소산 양자수율 (Φ_{NO})',
  ['콩: P1은 침수구가 대조구보다 높았고, P2는 건조·침수·복합구가 대조구보다 높았음','로메인: P1·P2 모두 처리 간 차이 없음'],
  '콩은 P1 침수구와 P2 모든 처리구의 ΦNO가 대조구보다 높고, 로메인은 처리 간 차이 없음',
  '(약 25초) 비조절 에너지 소산 양자수율 ΦNO입니다. 이 값이 높으면 조절되지 않는 형태로 에너지가 소산되어 광계 손상 위험이 커진다는 신호입니다. 콩은 P1에서 침수구가, P2에서 건조·침수·복합구가 모두 대조구보다 높았고, 로메인은 처리 간 차이가 없었습니다.'),
 ('PhiNPQ','ts_phinpq.png','LI-600 : 조절 에너지 소산 양자수율 (Φ_{NPQ})',
  ['콩: P1은 침수구가 대조구보다 높았고, P2는 건조구가 대조구·침수구보다 높았음','로메인: P1·P2 모두 처리 간 차이 없음'],
  '콩은 건조구에서 ΦNPQ가 높아 조절된 열소산이 증가, 로메인은 처리 간 차이 없음',
  '(약 25초) 조절된 에너지 소산 양자수율 ΦNPQ입니다. 콩은 P1에서 침수구가 대조구보다 높았고, P2에서는 건조구가 대조구와 침수구보다 높았습니다. 앞서 건조구에서 ΦPSII가 낮았던 것과 같이 보면, 흡수한 에너지가 광화학이 아니라 열소산 쪽으로 더 많이 쓰였다고 해석됩니다. 로메인은 처리 간 차이가 없었습니다.')]
FOOT='NPQ=(Fm−Fm′)/Fm′, Φ_{NO}=1/(NPQ+1+qL(Fm/Fo−1)), Φ_{NPQ}=1−Φ_{PSII}−Φ_{NO} (Kramer et al., 2004; Oxborough & Baker, 1997). Fo·Fm: 가장 가까운 날짜의 암적응 처리별 평균. 표·그림: PAR·기온 공변량 보정 평균, 혼합모형(처리 고정, 측정일 임의효과; Ali et al., 2025; Gao et al., 2024), Tukey 문자.'
for s,(ind,fig,ttl,txt,cc,nt) in zip(slides,SPEC):
    set_para(s.shapes.title.text_frame.paragraphs[0],ttl)
    pic=[x for x in s.shapes if x.name=='Picture 13'][0]
    replace_picture(s,pic,Image.open('an/figs2/'+fig).convert('RGB'),valign='top')
    shp=[x for x in s.shapes if getattr(x,'has_table',False) and x.has_table][0]; t=shp.table
    set_cell(t.cell(0,0),'Treatment')
    for r_i,tr in enumerate(['ETc','ETc50','침수ETc','침수ETc50'],start=1):
        for c_i,(cr,pe) in enumerate([('bean','P1'),('bean','P2'),('lettuce','P1'),('lettuce','P2')],start=1): set_cell(t.cell(r_i,c_i),pv(cr,pe,ind,tr))
    set_paras2(byname(s,'TextBox 10'),[(a.split(':')[0]+':',a.split(':',1)[1]) for a in txt])
    conc(s,cc); foot(s,FOOT); s.notes_slide.notes_text_frame.text=nt
    compact_table(t); bottom_border(t)
    for e in s._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr'): pass
# unique shape ids per slide (copied slides keep ids; fine) ; fix timing
for s in prs.slides: fix_timing(s)
prs.save('intro/NPQ_슬라이드.pptx'); print('saved',len(prs.slides))
