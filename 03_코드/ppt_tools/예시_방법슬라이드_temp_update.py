# -*- coding: utf-8 -*-
"""v6: 선형 PAR 보정에 챔버 기온(Tair = LI-600 Tref) 공변량 추가 (code6). 표·문장·방법 슬라이드·그림 갱신."""
import glob, json, re, shutil, sys, copy, warnings
warnings.filterwarnings("ignore")
import numpy as np, pandas as pd
BS = chr(92)
W = BS.join(["C:", "Users", "yunjm", "pptwork", "v6"])
sys.path.insert(0, W); sys.path.insert(0, BS.join(["C:", "Users", "yunjm", "pptwork", "cmp"])); sys.path.insert(0, BS.join(["C:", "Users", "yunjm", "pptwork", "oral"]))
import oral_style as OS
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import statsmodels.formula.api as smf
C6 = W + BS + "code6"
sys.path.insert(0, C6)
import config as cfg6
from pptx import Presentation
from pptx.oxml.ns import qn
from v6lib import *
import fmt_lib

D = glob.glob(r"C:\Users\yunjm\OneDrive*\바탕*\연구*\밭콩*\1차*\261106*")[0]
PATH = D + BS + "콩_vs_상추_비교_v6.pptx"
shutil.copy2(PATH, W + BS + "v6_before_temp.pptx")
p = Presentation(W + BS + "v6_before_temp.pptx")
ORDER = ["ETc", "ETc50", "침수ETc", "침수ETc50"]
ps = pd.read_csv(C6 + BS + "intermediate" + BS + "phase_stats.csv")
ps5 = pd.read_csv(W + BS + "code5" + BS + "intermediate" + BS + "phase_stats.csv")
rb = pd.read_csv(C6 + BS + "intermediate" + BS + "phase_robust.csv")
PL = json.load(open(W + BS + "par_phi_letters6.json", encoding="utf-8"))
RG = json.load(open(W + BS + "par_gsw_restored6.json", encoding="utf-8"))
fp = lambda v: "<0.001" if v < 0.001 else f"{v:.3f}"


def find(part):
    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if sh.shape_id == 2 and sh.has_text_frame and part in sh.text_frame.text:
                return i, s
    raise KeyError(part)


def tb(s):
    return [x for x in s.shapes if getattr(x, "has_table", False) and x.has_table][0].table


def rrep(slide, sid, old, new):
    n = 0
    for pa in [x for x in slide.shapes if x.shape_id == sid][0].text_frame.paragraphs:
        for r in pa.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    if n != 1:
        print("  !! replace count", n, old)
    return n


# ================================================================ 1. 문구: "... Q_amb" 뒤에 기온(T_air) 추가 (방법 슬라이드 7쪽 제외)
def add_temp(pa):
    runs = pa.runs
    for i in range(len(runs) - 2):
        if re.search(r"PAR[ (]Q$", runs[i].text) and runs[i + 1].text == "amb" and runs[i + 2].text.startswith(")"):
            close, rest = runs[i + 2], runs[i + 2].text[1:]
            sep = "와 " if rest.startswith("로 ") else (" + " if (rest.startswith(" 모형") or rest.startswith(" + 측정일")) else "·")
            r_t = copy.deepcopy(close._r); r_a = copy.deepcopy(runs[i + 1]._r); r_c = copy.deepcopy(close._r)
            close.text = ")"
            r_t.find(qn("a:t")).text = sep + "기온(T"; r_a.find(qn("a:t")).text = "air"; r_c.find(qn("a:t")).text = ")" + rest
            close._r.addnext(r_t); r_t.addnext(r_a); r_a.addnext(r_c)
            return True
    return False


n_add = 0
for k, s in enumerate(p.slides, 1):
    if k == 7:
        continue
    for sh in s.shapes:
        if sh.has_text_frame:
            for pa in sh.text_frame.paragraphs:
                n_add += add_temp(pa)
print("기온 문구 추가:", n_add)

# ================================================================ 2. Fs, ETR 표
def adj_table(s, ind, dec):
    t = tb(s)
    cols = {1: ("bean", "P1"), 2: ("bean", "P2"), 3: ("lettuce", "P1"), 4: ("lettuce", "P2")}
    for c, (crop, per) in cols.items():
        q = ps[(ps["지표"] == ind) & (ps["작물"] == crop) & (ps["국면"] == per)].set_index("처리")
        for r, tr in enumerate(ORDER, start=1):
            cell_set(t, r, c, f"{q.loc[tr, '보정평균']:.{dec}f} {q.loc[tr, '문자']}" if tr in q.index else "–")
        wi = q["상호작용P"].dropna()
        cell_set(t, 5, c, "–" if wi.empty else fp(wi.iloc[0]))
        cell_set(t, 6, c, fp(q["처리P"].iloc[0]))


q = lambda ind, crop, per: ps[(ps["지표"] == ind) & (ps["작물"] == crop) & (ps["국면"] == per)].set_index("처리")
i_fs, s_fs = find("정상 상태 형광"); adj_table(s_fs, "fs", 0)
i_etr, s_etr = find("전자전달속도"); adj_table(s_etr, "etr", 0)
f1, f2 = q("fs", "bean", "P1"), q("fs", "bean", "P2")
oth = [f2.loc[t, "보정평균"] for t in ("ETc50", "침수ETc", "침수ETc50")]
rrep(s_fs, 11, "131~133", f"{min(oth):.0f}~{max(oth):.0f}")
rrep(s_fs, 11, "대조구 114", f"대조구 {f2.loc['ETc', '보정평균']:.0f}")
rrep(s_fs, 11, "대조구 130 (b)", f"대조구 {f1.loc['ETc', '보정평균']:.0f} (b)")
rrep(s_fs, 11, "침수 139 (a)", f"침수 {f1.loc['침수ETc', '보정평균']:.0f} (a)")
eb1, eb2, el1, el2 = q("etr", "bean", "P1"), q("etr", "bean", "P2"), q("etr", "lettuce", "P1"), q("etr", "lettuce", "P2")
paras = [x for x in s_etr.shapes if x.shape_id == 11][0].text_frame.paragraphs
wi2 = eb2["상호작용P"].dropna().iloc[0]
l1 = eb1["문자"]
paras[1].runs[0].text = (f"콩 P1 대조구 {eb1.loc['ETc', '보정평균']:.0f} ({l1['ETc']}) vs 침수 {eb1.loc['침수ETc', '보정평균']:.0f} ({l1['침수ETc']}) (P={eb1['처리P'].iloc[0]:.3f}, 광도·기온 보정 방식에 따라 유의성이 달라짐), "
                         f"P2 대조구 {eb2.loc['ETc', '보정평균']:.0f} (a) > 복합 {eb2.loc['침수ETc50', '보정평균']:.0f} (bc) > 건조 {eb2.loc['ETc50', '보정평균']:.0f} (c) (W × I P={wi2:.3f})")
for r in paras[1].runs[1:]:
    r.text = ""
paras[2].runs[0].text = f"상추 P1 {el1['보정평균'].min():.0f}~{el1['보정평균'].max():.0f} (P={el1['처리P'].iloc[0]:.2f}), P2 {el2['보정평균'].min():.0f}~{el2['보정평균'].max():.0f} (P={el2['처리P'].iloc[0]:.2f}): 유의차 없음"
for r in paras[2].runs[1:]:
    r.text = ""
print("ETR P2 letters", eb2["문자"].to_dict())

# ================================================================ 3. ΦPSII 슬라이드
i_phi, s_phi = find("PSII: 콩 건조 저하")
t = tb(s_phi)
ph1 = q("phi", "bean", "P1")
cell_set(t, 1, 1, f"0.694 {ph1.loc['ETc', '문자']}"); cell_set(t, 3, 1, f"0.687 {ph1.loc['침수ETc', '문자']}")
for c, key in ((2, "300"), (3, "1200")):
    pr, lt = PL[key]["pred"], PL[key]["letters"]
    for r, tr in enumerate(ORDER, start=1):
        cell_set(t, r, c, f"{pr[tr]:.2f} {lt[tr]}")
rbp = rb[rb["지표"] == "phi"].drop_duplicates(["작물", "국면"]).set_index(["작물", "국면"])
cell_set(t, 6, 1, f"{rbp.loc[('bean', 'P1'), '처리x광도P']:.2f}"); cell_set(t, 6, 2, fp(rbp.loc[("bean", "P2"), "처리x광도P"])); cell_set(t, 6, 3, fp(rbp.loc[("bean", "P2"), "처리x광도P"]))
cell_set(t, 6, 4, f"{rbp.loc[('lettuce', 'P1'), '처리x광도P']:.2f}"); cell_set(t, 6, 5, f"{rbp.loc[('lettuce', 'P2'), '처리x광도P']:.2f}")
pic = [x for x in s_phi.shapes if x.shape_type == 13 and x.name != "Logo"][0]
swap_pic(s_phi, pic.shape_id, C6 + BS + "figures" + BS + "phi_light.png", width_in=7.55, top_in=1.3)
pa = [x for x in s_phi.shapes if x.shape_id == 11][0].text_frame.paragraphs
c1, c2 = ph1.loc["ETc", "보정평균"], ph1.loc["침수ETc", "보정평균"]
pa[0].runs[0].text = f"P1: 측정일 반복 검정은 두 작물 모두 차이 없음 (콩 P=0.44). 광도·기온 보정 평균은 침수 {c2:.3f} < 대조구 {c1:.3f} (P<0.001)로 차이가 작고 보정 방식에 따라 달라짐"
for r in pa[0].runs[1:]:
    r.text = ""
p3, p12 = PL["300"]["pred"], PL["1200"]["pred"]
pa[1].runs[-1].text = f" 저광(300)에서는 {min(p3.values()):.2f}~{max(p3.values()):.2f}로 차이 없음, 고광(1,200)에서 건조 {p12['ETc50']:.2f} < 복합 {p12['침수ETc50']:.2f} < 침수 {p12['침수ETc']:.2f} < 대조구 {p12['ETc']:.2f}"
n = 0
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for pp in sh.text_frame.paragraphs:
                for r in pp.runs:
                    if "0.69 → 0.50" in r.text:
                        r.text = r.text.replace("0.69 → 0.50", f"{p12['ETc']:.2f} → {p12['ETc50']:.2f}"); n += 1
print("결론 수치 교체:", n)

# ================================================================ 4. 부록 gsw 복원
i_ap, s_ap = find("gsw 전체 측정값 포함")
old_wi = ps[(ps["지표"] == "gsw") & (ps["작물"] == "bean") & (ps["국면"] == "P2")]["상호작용P"].dropna().iloc[0]
new_wi = RG["P2"][1]["상호작용P"]
cell_set(tb(s_ap), 6, 2, f"{old_wi:.2f} → {new_wi:.3f}")
for pp in [x for x in s_ap.shapes if x.shape_id == 11][0].text_frame.paragraphs:
    for r in pp.runs:
        if re.search(r"0\.\d+ → 0\.\d{3}", r.text):
            r.text = re.sub(r"0\.\d+ → 0\.\d{3}", f"{old_wi:.2f} → {new_wi:.3f}", r.text)
print("gsw 침수×관수 P:", old_wi, new_wi)

# ================================================================ 5. 방법 슬라이드 (7쪽): 그림 + 표 + 문장
x = cfg6.load_li600(); x = x[x.Qamb > 0].copy()
b = x[x.작물 == "bean"]
P2 = cfg6.common_days(b[b.DAT > 9]).dropna(subset=["PhiPS2", "Tair"]).copy()
m = smf.ols("PhiPS2 ~ C(처리) + C(날짜) + Qamb + Tair", P2).fit()
beta, bT, pT = m.params["Qamb"], m.params["Tair"], m.pvalues["Tair"]
r_qt = b[["Qamb", "Tair"]].corr().iloc[0, 1]
OS.setup(10.5)
fig, axs = plt.subplots(1, 3, figsize=(11.6, 3.5), gridspec_kw={"width_ratios": [1.25, 1, 1]})
ax = axs[0]
for t_ in ORDER:
    qq = P2[P2["처리"] == t_]
    ax.scatter(qq.Qamb, qq.PhiPS2, s=13, color=OS.col("bean", t_) if OS.FILLED[t_] else "white", edgecolor=OS.col("bean", t_), marker=OS.MK[t_], lw=0.9, alpha=0.8)
xx = np.linspace(P2.Qamb.min(), P2.Qamb.max(), 50)
ax.plot(xx, P2.PhiPS2.mean() + beta * (xx - P2.Qamb.mean()), color="#C00000", lw=2.0)
ax.text(0.97, 0.95, "common slope (T$_{air}$ held)\nβ = %.3f per 100 µmol m$^{-2}$ s$^{-1}$" % (beta * 100), transform=ax.transAxes, ha="right", va="top", fontsize=8.5, color="#C00000")
ax.set_xlabel("Q$_{amb}$ (PAR, µmol m$^{-2}$ s$^{-1}$)"); ax.set_ylabel(r"$\Phi_{PSII}$")
ax.set_title("A. Soybean P2: $\\Phi_{PSII}$ vs PAR", fontsize=10, loc="left")
dates = ["2026-08-13", "2026-08-18", "2026-09-11", "2026-10-02"]
for ax, col, ttl, yl in ((axs[1], "Qamb", "B. Same-day PAR by treatment", "Q$_{amb}$ (µmol m$^{-2}$ s$^{-1}$)"), (axs[2], "Tair", "C. Same-day air temp. by treatment", "T$_{air}$ (°C)")):
    for t_ in ORDER:
        mu, se, xs = [], [], []
        for i, d in enumerate(dates):
            qq = b[(b.날짜.astype(str).str[:10] == d) & (b.처리 == t_)][col]
            if len(qq):
                mu.append(qq.mean()); se.append(qq.sem()); xs.append(i + (ORDER.index(t_) - 1.5) * 0.14)
        ax.errorbar(xs, mu, se, fmt=OS.MK[t_], color=OS.col("bean", t_), mfc=OS.col("bean", t_) if OS.FILLED[t_] else "white", ms=6.5, capsize=2, lw=1.2)
    ax.set_xticks(range(4)); ax.set_xticklabels(["8/13", "8/18", "9/11", "10/2"]); ax.set_ylabel(yl); ax.set_title(ttl, fontsize=10, loc="left")
OS.top_legend(fig, line=False, pre=False, crops=("bean",), y=1.0, fontsize=9)
fig.tight_layout(rect=[0, 0, 1, 0.9]); fig.savefig(W + BS + "par_method6.png", dpi=220); plt.close(fig)
print("beta/100:", beta * 100, "bT:", bT, "pT:", pT, "r(Q,T):", r_qt)

i7, s7 = find("LI-600 광도 보정")
set_title(s7, "LI-600 보정: 선형 PAR(Q_{amb}) + 기온(T_{air})")
swap_pic(s7, [z for z in s7.shapes if z.shape_type == 13 and z.name != "Logo"][0].shape_id, W + BS + "par_method6.png", width_in=7.55, top_in=1.3)
gf = [z for z in s7.shapes if getattr(z, "has_table", False) and z.has_table][0]
t = gf.table
old = pd.read_csv(glob.glob(r"C:\Users\yunjm\OneDrive*\바탕*\연구*\밭콩*\1차*\cowork_분석자료\그래프_코드\intermediate\phase_stats.csv")[0])


def P_of(df, ind, per):
    return df[(df.지표 == ind) & (df.작물 == "bean") & (df.국면 == per)]["처리P"].iloc[0]


def sg(v):
    return "있음" if v < 0.05 else "없음"


rows = [["콩 P1 (대조구 vs 침수)", "log PAR P", "선형 PAR P", "+ 기온 P", "결론"]]
for ind, lab, per in (("phi", "ΦPSII", "P1"), ("etr", "ETR", "P1"), ("fs", "Fs", "P1"), ("gsw", "gsw", "P1"), ("phi", "ΦPSII 콩 P2 (4처리)", "P2"), ("etr", "ETR 콩 P2 (4처리)", "P2")):
    a_, b_, c_ = P_of(old, ind, per), P_of(ps5, ind, per), P_of(ps, ind, per)
    sig = [a_ < 0.05, b_ < 0.05, c_ < 0.05]
    concl = "모두 차이 있음" if all(sig) else ("모두 차이 없음" if not any(sig) else "차이 " + " → ".join(sg(v) for v in (a_, b_, c_)))
    rows.append([lab, fp(a_), fp(b_), fp(c_), concl])
for r_, rw in enumerate(rows):
    for c_, v in enumerate(rw):
        cell_set(t, r_, c_, v)
set_widths(gf, [2.0, 1.1, 1.2, 1.1, 2.1]); set_row_h(gf, 0.27)
print(rows)
dd = b[b.날짜.astype(str).str[:10] == "2026-10-02"].groupby("처리").Tair.mean()
set_text(shp(s7, 11), [
    "이유: 콩은 처리별로 묶어 측정해 같은 날에도 처리 간 광도·기온이 달랐음 (8/13 PAR 대조 970 vs 침수 632 µmol m^{-2} s^{-1}, 기온 33.3 vs 34.0 °C)",
    "모형: y = μ + 처리 + 측정일 + β_{1}·Q_{amb} + β_{2}·T_{air} + ε (T_{air} = LI-600 챔버 기온 Tref). Ali et al. 2025, Gao et al. 2024도 광도·온도를 공변량으로 보정",
    "보정 평균 = 균등 가중 LS mean (Q_{amb}, T_{air} 전체 평균에서 평가), Wald F, Tukey-Kramer 문자. 4처리 구간은 침수 × 관수 + 측정일 + Q_{amb} + T_{air}",
    f"효과 크기: 콩 P2 Φ_{{PSII}}는 PAR 100당 {beta * 100:.3f}, 기온은 유의하지 않음 (P={pT:.2f}). 처리 간 기온 차는 같은 날 평균 약 1.4 °C",
    f"한계: 선형은 관측 범위(PAR 약 100~1,900) 근사, PAR·기온 상관 r={r_qt:.2f}. 잎 온도(T_{{leaf}})는 처리의 영향을 받아 공변량에서 제외"])
set_text(shp(s7, 12), ["큰 차이(g_{sw}, F_{s}, 콩 P2)는 보정에 안정적 — 작은 차이(콩 P1 Φ_{PSII}, ETR)는 보정 방식에 민감"])
set_text(shp(s7, 13), ["그림 A: 콩 P2 개별 잎, 선 = 처리·측정일·기온 보정 후 PAR 공통 기울기. B·C: 측정일별 처리 평균 ± SE. 표: 콩 P1 대조구 vs 침수(P2는 4처리), P = 처리 효과(Wald F). Ali et al. (2025) J Exp Bot 76:4141; Gao et al. (2024) J Exp Bot 75:4904."])

# ================================================================ 6. 지표 첨자 변환 (새로 생긴 run)
fmt_lib.TOKENS.append(("sub", re.compile(fmt_lib.NA + r"Q(amb)" + fmt_lib.NB)))
ch = 0
for sl, pp in fmt_lib.all_paragraphs(p):
    for r in list(pp.findall(qn("a:r"))):
        ch += 1 if fmt_lib.convert_run(r) else 0
print("변환 run:", ch)
p.save(PATH)
print("saved")
