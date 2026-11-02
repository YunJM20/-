# -*- coding: utf-8 -*-
"""표를 '보정 평균(선형 PAR + 기온) 문자 (PAR만 보정; 원 평균)' 형식으로 통일. 상추 포함."""
import glob, json, re, sys, copy, shutil
import pandas as pd
BS = chr(92)
W = BS.join(["C:", "Users", "yunjm", "pptwork", "v6"])
sys.path.insert(0, W); sys.path.insert(0, "C:/Users/yunjm/pptwork/oral")
from pptx import Presentation
from pptx.oxml.ns import qn
from v6lib import *
import fmt_lib

D = glob.glob(r"C:\Users\yunjm\OneDrive*\바탕*\연구*\밭콩*\1차*\261106*")[0]
F = D + BS + "콩_vs_상추_비교_v6_보정설명추가.pptx"
shutil.copy2(F, W + BS + "v6_before_adjtab.pptx")
p = Presentation(W + BS + "v6_before_adjtab.pptx")
ORDER = ["ETc", "ETc50", "침수ETc", "침수ETc50"]
ps = pd.read_csv(W + r"\code6\intermediate\phase_stats.csv"); p5 = pd.read_csv(W + r"\code5\intermediate\phase_stats.csv")
PL5 = json.load(open(W + r"\par_phi_letters.json", encoding="utf-8")); PL6 = json.load(open(W + r"\par_phi_letters6.json", encoding="utf-8"))
EX = json.load(open(W + r"\extra_adj.json", encoding="utf-8"))
fp = lambda v: "<0.001" if v < 0.001 else f"{v:.3f}"


def find(part):
    for s in p.slides:
        for sh in s.shapes:
            if sh.shape_id == 2 and sh.has_text_frame and part in sh.text_frame.text:
                return s


def tb(s):
    return [x for x in s.shapes if getattr(x, "has_table", False) and x.has_table][0].table


def cell2(t, r, c, main, paren, sz=800):
    para = t.cell(r, c).text_frame.paragraphs[0]
    runs = para.runs
    runs[0].text = main
    for x in runs[1:]:
        x.text = ""
    new = copy.deepcopy(runs[0]._r)
    new.find(qn("a:t")).text = f" ({paren})"
    rpr = new.find(qn("a:rPr"))
    rpr.set("sz", str(sz))
    runs[0]._r.addnext(new)
    # 뒤쪽 빈 run 제거
    for x in para.runs[2:]:
        if not x.text:
            x._r.getparent().remove(x._r)


def rrep(s, sid, old, new):
    n = 0
    for pa in shp(s, sid).text_frame.paragraphs:
        for r in pa.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    print("  rrep", n, old[:24])


def get(df, ind, crop, per):
    return df[(df["지표"] == ind) & (df["작물"] == crop) & (df["국면"] == per)].set_index("처리")


def fill(t, ind, dec, cols, rows_skip=()):
    for c, (crop, per) in cols.items():
        a, b = get(ps, ind, crop, per), get(p5, ind, crop, per)
        for r, tr in enumerate(ORDER, start=1):
            if tr not in a.index:
                continue
            cell2(t, r, c, f"{a.loc[tr, '보정평균']:.{dec}f} {a.loc[tr, '문자']}", f"{b.loc[tr, '보정평균']:.{dec}f}; {a.loc[tr, '원평균']:.{dec}f}")


C4 = {1: ("bean", "P1"), 2: ("bean", "P2"), 3: ("lettuce", "P1"), 4: ("lettuce", "P2")}
KEY = "괄호 = (PAR만 보정; 원 평균)"

# ---- 9쪽 g_sw
s9 = find("기공전도도 (g")
fill(tb(s9), "gsw", 3, C4)
a1 = get(ps, "gsw", "bean", "P1")
rrep(s9, 11, "침수 0.27 vs 대조구 0.36", f"침수 {a1.loc['침수ETc', '보정평균']:.2f} vs 대조구 {a1.loc['ETc', '보정평균']:.2f}")
rrep(s9, 13, "표: 평균은 원 평균(보정 없음), 문자는 측정일·선형 PAR·기온 보정 모형의 Tukey.", "표: 측정일·선형 PAR·기온 보정 평균, Tukey 문자. " + KEY + ".")

# ---- 12쪽 Φ_PSII
s12 = find("PSII: 콩 건조 저하")
t = tb(s12)
fill(t, "phi", 3, {1: ("bean", "P1"), 4: ("lettuce", "P1"), 5: ("lettuce", "P2")})
for c, key in ((2, "300"), (3, "1200")):
    for r, tr in enumerate(ORDER, start=1):
        cell2(t, r, c, f"{PL6[key]['pred'][tr]:.2f} {PL6[key]['letters'][tr]}", f"{PL5[key]['pred'][tr]:.2f}")
l12 = [get(ps, "phi", "lettuce", per)["보정평균"] for per in ("P1", "P2")]
rrep(s12, 11, "0.59~0.70", f"{min(l12[0].min(), l12[1].min()):.2f}~{max(l12[0].max(), l12[1].max()):.2f}")
fn = shp(s12, 13)
for pa in fn.text_frame.paragraphs[:1]:
    pass
set_text(fn, ["표: 측정일·선형 PAR·기온 보정 평균(콩 P2는 처리 × PAR 모형의 PPFD 300·1,200 예측, 9/4 기준). 괄호 = (PAR만 보정; 원 평균), P2 예측은 PAR만 보정. P(측정일 반복) = 처리별 측정일 평균 ~ 침수 × 관수 + 측정일. 문자: Tukey(상추는 n.s.라 모두 a)."])

# ---- 13·14쪽 F_s, ETR
for key, ind in (("정상 상태 형광", "fs"), ("전자전달속도", "etr")):
    s = find(key)
    fill(tb(s), ind, 0, C4)
    rrep(s, 13, "Tukey 문자.", "Tukey 문자, " + KEY + ".")

# ---- 15쪽 상추 신호 크기 (Fs, Fm', Fm'/Fs 보정 / Fm 암적응은 그대로)
s15 = find("광계 II 손상")
t = tb(s15)
for c, v, dec in ((1, "Fs", 0), (2, "Fm'", 0), (4, "ratio", 2)):
    d = EX[v]; ctl = d["ETc"]["adjT"]
    for r, tr in enumerate(ORDER, start=1):
        cell_set(t, r, c, f"{d[tr]['adjT']:.{dec}f} ({d[tr]['adjT'] / ctl * 100:.0f}%) {d[tr]['let']}")
cell_set(t, 5, 4, f"n.s. ({EX['ratio']['P'][0]:.2f})")
rrep(s15, 13, "표·문자는 보정 없는 개별 잎 값 기준.", "표: Fs·Fm′·Fm′/Fs는 측정일·선형 PAR·기온 보정 평균(상추 P1), Fm(암적응)은 보정 없음.")

# ---- 11쪽 상추 10/10
s11 = find("상추 침수 직후(10/10)")
t = tb(s11)
for row, key, dec in ((4, "d1010_gsw", 3), (5, "d1010_phi", 3)):
    d = EX[key]
    for c, tr in ((1, "ETc"), (2, "침수ETc")):
        cell2(t, row, c, f"{d[tr]['adjT']:.{dec}f} ± {d[tr]['seT']:.{dec}f}", f"{d[tr]['adjP']:.{dec}f}; {d[tr]['raw']:.{dec}f}")
    cell_set(t, row, 3, f"{d['침수ETc']['adjT'] / d['ETc']['adjT'] * 100:.0f}%")
    cell2(t, row, 4, fp(d["P"][0]), f"{fp(d['P'][1])}; {'<0.001' if key.endswith('gsw') else '0.370'}")
dg, dp = EX["d1010_gsw"], EX["d1010_phi"]
rrep(s11, 11, "29% (0.251 → 0.073, P<0.001)", f"{dg['침수ETc']['adjT'] / dg['ETc']['adjT'] * 100:.0f}% ({dg['ETc']['adjT']:.3f} → {dg['침수ETc']['adjT']:.3f}, P<0.001)")
rrep(s11, 11, "92% (P=0.37)", f"{dp['침수ETc']['adjT'] / dp['ETc']['adjT'] * 100:.0f}% (P={dp['P'][0]:.2f})")
rrep(s11, 11, "(gsw 29%)", f"(gsw {dg['침수ETc']['adjT'] / dg['ETc']['adjT'] * 100:.0f}%)")
rrep(s11, 12, "gsw는 29%로", f"gsw는 {dg['침수ETc']['adjT'] / dg['ETc']['adjT'] * 100:.0f}%로")
rrep(s11, 13, "평균 ± SE, Welch t검정, *P<0.05. 보정 없음.", "g_sw·Φ_PSII 행: 선형 PAR·기온 보정 평균 ± SE, P = Wald F, 괄호 = (PAR만 보정; 원 평균, P는 PAR만; Welch t). 그 외 행은 보정 없음.")

# ---- 7쪽 적용 범위
s7 = find("LI-600 보정")
set_text(shp(s7, 13), ["그림 A: 콩 P2 개별 잎, 선 = 보정 후 PAR 기울기. B·C: 측정일별 평균 ± SE. 표: 콩 P1(P2는 4처리), Wald F. 적용: 9·11~15쪽 표는 PAR + 기온 보정 평균(괄호 = PAR만; 원 평균), 그림·10쪽 비율·Fm(암적응)은 보정 없음. Ali et al. (2025) JXB 76:4141; Gao et al. (2024) JXB 75:4904."])

fmt_lib.TOKENS.append(("sub", re.compile(fmt_lib.NA + r"Q(amb)" + fmt_lib.NB)))
ch = 0
for sl, pp in fmt_lib.all_paragraphs(p):
    for r in list(pp.findall(qn("a:r"))):
        ch += 1 if fmt_lib.convert_run(r) else 0
print("변환 run", ch)
p.save(F)
print("saved")
