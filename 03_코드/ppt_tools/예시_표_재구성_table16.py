# -*- coding: utf-8 -*-
"""16쪽(A-Ci·A-Q) 표를 '행 = 처리, 열 = 작물×파라미터' 로 재구성해 처리별 평균과 문자를 모두 넣는다."""
import glob, sys, copy, re, shutil
import pandas as pd
BS = chr(92)
W = BS.join(["C:", "Users", "yunjm", "pptwork", "v6"])
sys.path.insert(0, W); sys.path.insert(0, "C:/Users/yunjm/pptwork/oral")
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.oxml.ns import qn
from v6lib import *
import fmt_lib

D = glob.glob(r"C:\Users\yunjm\OneDrive*\바탕*\연구*\밭콩*\1차*\261106*")[0]
F = D + BS + "콩_vs_상추_비교_v6.pptx"
shutil.copy2(F, W + BS + "v6_before_t16.pptx")
p = Presentation(W + BS + "v6_before_t16.pptx")
s = [x for x in p.slides if any(sh.shape_id == 2 and sh.has_text_frame and "광합성 능력 (9, 10월)" in sh.text_frame.text for sh in x.shapes)][0]
gf = [x for x in s.shapes if getattr(x, "has_table", False) and x.has_table][0]
t = gf.table
pre = {(r, c): t.cell(r, c).text_frame.text for r in range(1, 6) for c in (1, 3)}   # 처리 전 평균
st = pd.read_csv(W + r"\code6\intermediate\param_stats.csv")
PARAMS = [("Vcmax", 0), ("Jmax", 0), ("Rd", 2), ("Amax", 1), ("alpha", 3)]
LAB = {"Vcmax": "Vcmax", "Jmax": "Jmax", "Rd": "Rd", "Amax": "Amax", "alpha": "α"}
ORDER = ["ETc", "ETc50", "침수ETc", "침수ETc50"]
ROWN = ["Pre-treatment", "ETc 100%", "ETc 50%", "Flooding-ETc 100%", "Flooding-ETc 50%", "P (ANOVA)"]

tbl = gf.table._tbl
trs = tbl.findall(qn("a:tr"))
hdr, body, last = trs[0], trs[1], trs[-1]
tcH0, tcH = hdr.findall(qn("a:tc"))[0], hdr.findall(qn("a:tc"))[1]
tcB0, tcB = body.findall(qn("a:tc"))[0], body.findall(qn("a:tc"))[1]
tcL0, tcL = last.findall(qn("a:tc"))[0], last.findall(qn("a:tc"))[1]
grid = tbl.tblGrid
for gc in list(grid):
    grid.remove(gc)
gc0 = None
widths = [1.25] + [0.62] * 10
for w in widths:
    g = etree.SubElement(grid, qn("a:gridCol")); g.set("w", str(int(w * 914400)))
for tr in trs:
    tbl.remove(tr)


def mk_row(tmpl_tr, t0, t1, texts, h=0.24):
    tr = copy.deepcopy(tmpl_tr)
    for tc in tr.findall(qn("a:tc")):
        tr.remove(tc)
    for i, tx in enumerate(texts):
        tc = copy.deepcopy(t0 if i == 0 else t1)
        tr.append(tc)
    for e in tr.findall(qn("a:extLst")):
        tr.remove(e)
    tr.set("h", str(int(h * 914400)))
    return tr


# 본문 행
rows = []
r1 = mk_row(hdr, tcH0, tcH, [""] * 11); r2 = mk_row(hdr, tcH0, tcH, [""] * 11)
rows += [r1, r2]
for k, name in enumerate(ROWN):
    tmpl, a, b = (last, tcL0, tcL) if name.startswith("P") else (body, tcB0, tcB)
    rows.append(mk_row(tmpl, a, b, [""] * 11))
for tr in rows:
    tbl.append(tr)
gf.height = Emu(int(0.24 * 8 * 914400))
for r in t.rows:
    r.height = Emu(int(0.24 * 914400))
set_widths(gf, widths)


def put(r, c, text, sz=900):
    para = t.cell(r, c).text_frame.paragraphs[0]
    runs = para.runs
    if not runs:
        para.add_run().text = text
        runs = para.runs
    else:
        runs[0].text = text
        for x in runs[1:]:
            x.text = ""
    for x in para.runs:
        x.font.size = Emu(sz * 127)


# 헤더 병합
for c0, name in ((1, "Soybean"), (6, "Lettuce")):
    tc = t.cell(0, c0)._tc
    tc.set("gridSpan", "5")
    for k in range(1, 5):
        t.cell(0, c0 + k)._tc.set("hMerge", "1")
    put(0, c0, name, 1000)
put(0, 0, "", 900); put(1, 0, "Treatment", 900)
for ci, crop in enumerate(("bean", "lettuce")):
    for pi, (pn, dec) in enumerate(PARAMS):
        c = 1 + ci * 5 + pi
        put(1, c, LAB[pn], 900)
        q = st[(st.작물 == crop) & (st.파라미터 == pn)].set_index("처리")
        pre_txt = pre[(1 + pi, 1 if crop == "bean" else 3)]
        put(2, c, pre_txt, 900)
        for ri, tr_ in enumerate(ORDER):
            put(3 + ri, c, f"{q.loc[tr_, '평균']:.{dec}f} {q.loc[tr_, 'Tukey']}", 900)
        P = q["ANOVA_P"].iloc[0]
        put(7, c, "<0.001" if P < 0.001 else f"{P:.3f}", 900)
for ri, name in enumerate(ROWN):
    put(2 + ri, 0, name, 900)
gf.top = Emu(int(4.32 * 914400))
fmt_lib.TOKENS.append(("sub", re.compile(fmt_lib.NA + r"Q(amb)" + fmt_lib.NB)))
for sl, pp in fmt_lib.all_paragraphs(p):
    for r in list(pp.findall(qn("a:r"))):
        fmt_lib.convert_run(r)
for pa in shp(s, 13).text_frame.paragraphs:
    for r in pa.runs:
        r.text = r.text.replace("표 문자: 대조구/건조/침수/복합.", "표: 처리별 평균, Tukey 문자.")
p.save(W + BS + "v6_t16_out.pptx")
print("saved")
