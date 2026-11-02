# -*- coding: utf-8 -*-
"""비교 PPT v5 -> v6 편집 도우미 (python-pptx + lxml)"""
import copy, io, re
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu
from PIL import Image

R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def shp(slide, sid):
    return [x for x in slide.shapes if x.shape_id == sid][0]


def delete(slide, sid):
    el = shp(slide, sid)._element
    el.getparent().remove(el)


# ---------------------------------------------------------------- 텍스트
def copy_paras(dst, src):
    """src 도형의 단락(a:p) 전체를 dst 도형으로 복사 (서식 포함)"""
    d, s = dst.text_frame._txBody, src.text_frame._txBody
    for p in d.findall(qn("a:p")):
        d.remove(p)
    for p in s.findall(qn("a:p")):
        d.append(copy.deepcopy(p))


_TOK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def _runs_for(line, rpr_tmpl, bold=None):
    out = []
    for tok in _TOK.split(line):
        if not tok:
            continue
        r = etree.Element(qn("a:r"))
        rpr = copy.deepcopy(rpr_tmpl) if rpr_tmpl is not None else etree.Element(qn("a:rPr"))
        for att in ("baseline", "b"):
            if att in rpr.attrib:
                del rpr.attrib[att]
        if tok.startswith("_{"):
            rpr.set("baseline", "-25000"); tok = tok[2:-1]
        elif tok.startswith("^{"):
            rpr.set("baseline", "30000"); tok = tok[2:-1]
        if bold:
            rpr.set("b", "1")
        r.append(rpr)
        t = etree.SubElement(r, qn("a:t")); t.text = tok
        out.append(r)
    return out


def set_text(shape, lines, bold_idx=()):
    """단락별 텍스트 설정. 기존 단락 서식(pPr, 첫 run rPr)을 복제. line 안의 _{x} 아래첨자, ^{x} 위첨자.
    bold_idx: 굵게 처리할 줄 번호(0부터)"""
    tb = shape.text_frame._txBody
    ps = tb.findall(qn("a:p"))
    tmpl = ps[0]
    rpr0 = None
    for p in ps:
        r = p.find(qn("a:r"))
        if r is not None and r.find(qn("a:rPr")) is not None:
            rpr0 = r.find(qn("a:rPr")); break
    ppr0 = tmpl.find(qn("a:pPr"))
    for p in ps:
        tb.remove(p)
    for i, line in enumerate(lines):
        base = ps[i] if i < len(ps) else ps[-1]
        r_ = base.find(qn("a:r"))
        rpr = r_.find(qn("a:rPr")) if (r_ is not None and r_.find(qn("a:rPr")) is not None) else rpr0
        p = etree.SubElement(tb, qn("a:p"))
        ppr = base.find(qn("a:pPr"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        for r in _runs_for(line, rpr, bold=(i in bold_idx)):
            p.append(r)


def set_title(slide, text, sid=2):
    """제목 도형: 첫 run(영문 eyebrow)과 줄바꿈은 유지, 이후를 text로 교체 (_{ }, ^{ } 지원)"""
    sh = shp(slide, sid)
    p = sh.text_frame._txBody.find(qn("a:p"))
    kids = list(p)
    br = [k for k in kids if k.tag == qn("a:br")][0]
    after = kids[kids.index(br) + 1:]
    tmpl = [k for k in after if k.tag == qn("a:r")][0].find(qn("a:rPr"))
    for k in after:
        p.remove(k)
    for r in _runs_for(text, tmpl):
        p.append(r)


def set_eyebrow(slide, text, sid=2):
    p = shp(slide, sid).text_frame._txBody.find(qn("a:p"))
    p.find(qn("a:r")).find(qn("a:t")).text = text


# ---------------------------------------------------------------- 표
def cell_set(tbl, r, c, text):
    para = tbl.cell(r, c).text_frame.paragraphs[0]
    runs = para.runs
    if not runs:
        run = para.add_run(); run.text = text; return
    runs[0].text = text
    for x in runs[1:]:
        x.text = ""


def table_copy(dst_tbl, src_tbl, skip_first_col_rows=(1, 2, 3, 4), rows=None):
    nr = min(len(dst_tbl.rows), len(src_tbl.rows))
    nc = min(len(dst_tbl.columns), len(src_tbl.columns))
    for r in range(nr):
        for c in range(nc):
            if c == 0 and r in skip_first_col_rows:
                continue
            cell_set(dst_tbl, r, c, src_tbl.cell(r, c).text_frame.text)


def add_col(gf, after, widths):
    tbl = gf.table._tbl
    gcs = tbl.tblGrid.findall(qn("a:gridCol"))
    gcs[after].addnext(copy.deepcopy(gcs[after]))
    for tr in tbl.findall(qn("a:tr")):
        tcs = tr.findall(qn("a:tc"))
        tcs[after].addnext(copy.deepcopy(tcs[after]))
    set_widths(gf, widths)


def set_widths(gf, widths):
    for c, w in zip(gf.table.columns, widths):
        c.width = Emu(int(w * 914400))
    gf.width = Emu(int(sum(widths) * 914400))


def add_rows(gf, n):
    """끝에서 두 번째 행(경계선 없는 일반 행)을 복제해 마지막 행 앞에 삽입"""
    tbl = gf.table._tbl
    for _ in range(n):
        trs = tbl.findall(qn("a:tr"))
        trs[-2].addnext(copy.deepcopy(trs[-2]))


def del_row(gf, idx):
    tbl = gf.table._tbl
    tr = tbl.findall(qn("a:tr"))[idx]
    tbl.remove(tr)


def set_row_h(gf, h_in):
    for r in gf.table.rows:
        r.height = Emu(int(h_in * 914400))
    gf.height = Emu(int(h_in * 914400 * len(gf.table.rows)))


# ---------------------------------------------------------------- 그림
def swap_pic(slide, sid, path, width_in=None, top_in=None, left_in=None, max_h_in=None):
    sh = shp(slide, sid)
    ip, rid = slide.part.get_or_add_image_part(path)
    sh._element.blipFill.blip.set(R_EMBED, rid)
    w, h = Image.open(path).size
    wi = width_in if width_in else sh.width / 914400
    hi = wi * h / w
    if max_h_in and hi > max_h_in:
        hi = max_h_in; wi = hi * w / h
    cx = (sh.left + sh.width / 2) / 914400
    sh.width, sh.height = Emu(int(wi * 914400)), Emu(int(hi * 914400))
    sh.left = Emu(int((left_in if left_in is not None else cx - wi / 2) * 914400))
    if top_in is not None:
        sh.top = Emu(int(top_in * 914400))


# ---------------------------------------------------------------- 슬라이드 복제·정렬
def clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    sp_new = new.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        sp_new.append(copy.deepcopy(el))
    for blip in sp_new.iter(qn("a:blip")):
        rid = blip.get(R_EMBED)
        part = src.part.related_part(rid)
        _, nrid = new.part.get_or_add_image_part(io.BytesIO(part.blob))
        blip.set(R_EMBED, nrid)
    return new


def reorder(prs, order):
    """order: 현재 슬라이드 객체의 새 순서 리스트"""
    lst = prs.slides._sldIdLst
    ids = {s.slide_id: e for s, e in zip(prs.slides, list(lst))}
    elems = [ids[s.slide_id] for s in order]
    for e in list(lst):
        lst.remove(e)
    for e in elems:
        lst.append(e)
