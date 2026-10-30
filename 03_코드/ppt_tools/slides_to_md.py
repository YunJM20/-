# -*- coding: utf-8 -*-
from pptx import Presentation
def para_text(pa): return "".join(r.text for r in pa.runs).strip()
def slide_md(s, n):
    out = []
    title = [sh for sh in s.shapes if sh.shape_id == 2 and sh.has_text_frame]
    t = "";
    if title:
        ps = title[0].text_frame.paragraphs
        t = " — ".join(x for x in (para_text(p) for p in ps) if x)
        t = title[0].text_frame.text.replace("\n", " | ").replace("\x0b", " | ")
    out.append(f"### 슬라이드 {n}: {t}\n")
    items = []
    for sh in s.shapes:
        if sh.shape_id == 2: continue
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = [[c.text_frame.text.replace("\n", " ").replace("|", "/") for c in r.cells] for r in sh.table.rows]
            md = "| " + " | ".join(rows[0]) + " |\n|" + "|".join(["---"] * len(rows[0])) + "|\n" + "".join("| " + " | ".join(r) + " |\n" for r in rows[1:])
            items.append((sh.top, sh.left, "table", md))
        elif sh.has_text_frame:
            lines = [para_text(p) for p in sh.text_frame.paragraphs]; lines = [l for l in lines if l]
            if not lines or sh.name in ("Logo",): continue
            if sh.name.startswith("슬라이드 번호"): continue
            kind = "foot" if sh.name.startswith("Footnote") else ("concl" if sh.name.startswith("Conclusion") else ("bul" if sh.name.startswith("TextBox") else "txt"))
            items.append((sh.top, sh.left, kind, lines))
    for top, left, kind, c in sorted(items, key=lambda x: (x[0] // 300000, x[1])):
        if kind == "table": out.append(c)
        elif kind == "bul" or kind == "txt": out += [f"- {l}" for l in c]
        elif kind == "concl": out.append("\n**결론 박스**: " + " ".join(c))
        elif kind == "foot": out.append("\n*각주*: " + " ".join(c))
    return "\n".join(out) + "\n"
if __name__ == "__main__":
    p = Presentation("cur2.pptx")
    open("slides_dump.md", "w", encoding="utf-8").write("\n".join(slide_md(s, i) for i, s in enumerate(p.slides, 1)))
