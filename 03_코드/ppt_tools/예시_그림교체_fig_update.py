# -*- coding: utf-8 -*-
"""그래프를 PAR·기온 보정 일평균(code7)으로 교체, 관련 문장·각주 갱신, 15쪽 표 괄호·박스 조정. 보정 전 그래프(code6 이전 산출물)는 건드리지 않는다."""
import glob, json, re, sys, copy, shutil
import pandas as pd
from pptx.util import Emu, Pt
BS = chr(92)
W = BS.join(["C:", "Users", "yunjm", "pptwork", "v6"])
sys.path.insert(0, W); sys.path.insert(0, "C:/Users/yunjm/pptwork/oral")
from pptx import Presentation
from pptx.oxml.ns import qn
from v6lib import *
import fmt_lib

D = glob.glob(r"C:\Users\yunjm\OneDrive*\바탕*\연구*\밭콩*\1차*\261106*")[0]
F = D + BS + "콩_vs_상추_비교_v6.pptx"
shutil.copy2(F, W + BS + "v6_before_figs.pptx")
p = Presentation(W + BS + "v6_before_figs.pptx")
EX = json.load(open(W + BS + "extra_adj.json", encoding="utf-8"))
ORDER = ["ETc", "ETc50", "침수ETc", "침수ETc50"]
FIG7 = W + BS + "code7" + BS + "figures" + BS


def find(part):
    for s in p.slides:
        for sh in s.shapes:
            if sh.shape_id == 2 and sh.has_text_frame and part in sh.text_frame.text:
                return s


def rrep(s, sid, old, new):
    n = 0
    for pa in shp(s, sid).text_frame.paragraphs:
        for r in pa.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); n += 1
    print("  rrep", sid, n, old[:26])


def pic(s, path):
    pc = [z for z in s.shapes if z.shape_type == 13 and z.name != "Logo"][0]
    swap_pic(s, pc.shape_id, path, width_in=pc.width / 914400, top_in=pc.top / 914400)


def tbl(s):
    return [x for x in s.shapes if getattr(x, "has_table", False) and x.has_table][0]


# ---------------- 그림 교체
pairs = (("기공전도도 (g", "ts_gsw.png"), ("기공 회복 속도", "recovery.png"), ("PSII: 콩 건조 저하", "phi_light.png"), ("정상 상태 형광", "ts_fs.png"),
         ("전자전달속도", "ts_etr.png"), ("광계 II 손상", "fvfm_signal.png"))
for key, fn in pairs:
    pic(find(key), FIG7 + fn)
pic(find("상추 침수 직후"), W + BS + "lettuce_1010_adj.png")

# ---------------- 비율·수치 (보정 일평균)
d = pd.read_csv(W + r"\code7\intermediate\daily_stats.csv"); a = d[(d.지표 == "gsw")]
g = {c: a[a.작물 == c].pivot_table(index="DAT", columns="처리", values="평균") for c in ("bean", "lettuce")}
r = {c: (g[c]["침수ETc"] / g[c]["ETc"] * 100) for c in g}
lmin = r["lettuce"].dropna().min(); lday = int(r["lettuce"].idxmin())
b29 = r["bean"].loc[2:9]
print("lettuce min", lday, round(lmin), "bean 2-9", round(b29.min()), round(b29.max()), "d9", round(r["bean"][9]), "d17", round(r["bean"][17]), "d23", round(r["bean"][23]))
s10 = find("기공 회복 속도")
rrep(s10, 11, "(29%)", f"({lmin:.0f}%)")
rrep(s10, 11, "39~80%", f"{b29.min():.0f}~{b29.max():.0f}%")
rrep(s10, 11, "DAT 9 48%", f"DAT 9 {r['bean'][9]:.0f}%")
rrep(s10, 11, "DAT 23 105%", f"DAT 23 {r['bean'][23]:.0f}%")
rrep(s10, 16, "DAT 9: 48%", f"DAT 9: {r['bean'][9]:.0f}%")
rrep(s10, 13, "비율은 보정 없는 당일 평균 g", "비율은 PAR·기온 보정 일평균 g")
s9 = find("기공전도도 (g")
rrep(s9, 11, "29% (0.07 vs 0.25)", f"{lmin:.0f}% ({g['lettuce']['침수ETc'][lday]:.2f} vs {g['lettuce']['ETc'][lday]:.2f})")
rrep(s9, 13, "그림(보정 없음):", "그림(보정 일평균):")
for key in ("정상 상태 형광", "전자전달속도"):
    rrep(find(key), 13, "그림(무보정 일평균):", "그림(보정 일평균):")

# ---------------- 7쪽
s7 = find("LI-600 보정")
set_text(shp(s7, 13), ["그림 A: 콩 P2 개별 잎, 선 = 보정 후 PAR 기울기. B·C: 측정일별 평균 ± SE. 표: 콩 P1(P2는 4처리), Wald F. 적용: 9·11~15쪽 표·그림과 10쪽 비율은 PAR + 기온 보정(표 괄호 = PAR만; 원 평균), 처리 전(X)·암적응 F_{m}·F_{v}/F_{m}은 보정 없음. Ali et al. (2025) JXB 76:4141; Gao et al. (2024) JXB 75:4904."])
set_text(shp(s7, 11), [
 "이유: 같은 날에도 처리 간 PAR·기온이 달랐음 (8/13 PAR 970 vs 632 µmol m^{-2} s^{-1}, 기온 33.3 vs 34.0 °C)",
 "모형: y = μ + 처리 + 측정일 + β_{1}·Q_{amb} + β_{2}·T_{air} + ε. Ali 2025, Gao 2024도 광도·온도를 공변량으로 보정",
 "보정 평균: 균등 가중 LS mean, Wald F, Tukey-Kramer. 그림: 개체값을 전체 평균 Q_{amb}·T_{air}로 보정한 일평균",
 "콩 P2 Φ_{PSII}: PAR 100당 −0.022, 기온 효과 없음 (P=0.68)",
 "한계: 선형 근사(PAR 약 100~1,900), PAR·기온 r=0.29, T_{air} = LI-600 챔버 기온(Tref), 잎 온도는 처리 영향을 받아 제외"])

# ---------------- 11쪽 각주
s11 = find("상추 침수 직후")
set_text(shp(s11, 13), ["잎상추 10/10 A-Ci(각 3곡선), LI-600(상엽, 처리당 12잎, Qamb 대조 899 vs 침수 857), 암적응 Fv/Fm(9잎). gsw·ΦPSII 행·그림 패널 = 선형 PAR·기온 보정(P = Wald F; 괄호 = PAR만 보정; 원 평균, P는 PAR만; Welch t), 그 외는 보정 없음(Welch t)."])

# ---------------- 15쪽: 표 괄호, 박스 축소
s15 = find("광계 II 손상")
g15 = tbl(s15); t = g15.table


def cell2(t, rr_, cc, main, paren, sz=800):
    para = t.cell(rr_, cc).text_frame.paragraphs[0]
    runs = para.runs
    runs[0].text = main
    for x in runs[1:]:
        x.text = ""
    nw = copy.deepcopy(runs[0]._r); nw.find(qn("a:t")).text = f" ({paren})"; nw.find(qn("a:rPr")).set("sz", str(sz))
    runs[0]._r.addnext(nw)
    for x in para.runs[2:]:
        if not x.text:
            x._r.getparent().remove(x._r)


for c, v, dec in ((1, "Fs", 0), (2, "Fm'", 0), (4, "ratio", 2)):
    dd = EX[v]; ctl = dd["ETc"]["adjT"]
    for rr_, tr in enumerate(ORDER, start=1):
        cell2(t, rr_, c, f"{dd[tr]['adjT']:.{dec}f} ({dd[tr]['adjT'] / ctl * 100:.0f}%) {dd[tr]['let']}", f"{dd[tr]['adjP']:.{dec}f}; {dd[tr]['raw']:.{dec}f}")
set_widths(g15, [1.45, 1.8, 1.8, 1.5, 1.8])
for sid, (left, width) in {9: (9.05, 3.75), 10: (9.25, 3.35), 11: (9.25, 3.45)}.items():
    sh = shp(s15, sid); sh.left = Emu(int(left * 914400)); sh.width = Emu(int(width * 914400))
for pa_ in shp(s15, 11).text_frame.paragraphs:
    for rr in pa_.runs:
        rr.font.size = Pt(11)
set_text(shp(s15, 13), ["Fv/Fm: LI-600 야간 암적응, 평균 ± SE, * 당일 ANOVA P<0.05. 콩 건조·복합구는 DAT 29만(9/28 LI-6800 값 제외). 신호 크기 표·그림: 상추 P1(DAT 5~10) Fs·Fm′·Fm′/Fs는 PAR·기온 보정, 표 괄호 = (PAR만 보정; 원 평균), 암적응 Fm(10/12~10/17)은 보정 없음. 문자: Tukey."])

fmt_lib.TOKENS.append(("sub", re.compile(fmt_lib.NA + r"Q(amb)" + fmt_lib.NB)))
fmt_lib.TOKENS.append(("sub", re.compile(fmt_lib.NA + r"T(air)" + fmt_lib.NB)))
ch = 0
for sl, pp in fmt_lib.all_paragraphs(p):
    for rr in list(pp.findall(qn("a:r"))):
        ch += 1 if fmt_lib.convert_run(rr) else 0
print("변환 run", ch)
p.save(W + BS + "v6_figs_out.pptx")
print("saved local")
